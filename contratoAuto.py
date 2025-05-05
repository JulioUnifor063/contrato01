from PyQt6.QtWidgets import QComboBox
import locale
locale.setlocale(locale.LC_TIME, 'Portuguese_Brazil')
from datetime import datetime
import sys
import os
import re
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt
from PyQt6.QtWidgets import QApplication, QMainWindow, QVBoxLayout, QWidget, QLabel, QLineEdit, QPushButton, QDateEdit, QMessageBox, QCheckBox
from PyQt6.QtCore import QDate
from PyQt6.QtGui import QDoubleValidator, QIntValidator
from num2words import num2words
import pandas as pd


def capitalize_first_letter(text):
    return text.capitalize()

def obter_diretorio_base():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

def converter_xls_para_xlsx(arquivo_origem_xls, arquivo_destino_xlsx):
    df = pd.read_excel(arquivo_origem_xls, engine='xlrd')
    df.to_excel(arquivo_destino_xlsx, index=False, engine='openpyxl')
    print('Arquivo .xls convertido para .xlsx com sucesso!')
    os.remove(arquivo_origem_xls)
    print('Arquivo .xls original removido com sucesso!')

def replace_placeholder(paragraph, placeholder, replacement, font_size=Pt(9), to_upper=False):
    for run in paragraph.runs:
        if placeholder in run.text:
            if to_upper:
                replacement = replacement.upper()
            run.text = run.text.replace(placeholder, str(replacement))
            run.font.size = font_size


def formatar_valor_mensalidade(valor):
    try:
        valor_float = float(valor)
        valor_formatado = f"R$ {valor_float:,.2f}".replace('.', 'X').replace(',', '.').replace('X', ',')
        return valor_formatado
    except ValueError:
        return "R$ Erro"

def formatar_valor_adesao(valor, prefixo=""):
    try:
        valor_float = float(valor)
        valor_formatado = f"R$ {valor_float:,.2f}".replace('.', 'X').replace(',', '.').replace('X', ',')
        return f"{prefixo} {valor_formatado}"
    except ValueError:
        return "R$ Erro"

def atualizar_documento(cidade, data_contrato, data_inicio, data_fim, valor_mensalidade, qtd_acesso, valor_adesao):
    diretorio_base = obter_diretorio_base()

    arquivo_origem_xls = os.path.join(diretorio_base, 'Report.xls')
    arquivo_origem_xlsx = os.path.join(diretorio_base, 'Report.xlsx')
    
    if cidade == "Fortaleza":
        arquivo_word_origem = os.path.join(diretorio_base, 'Fortaleza.docx')
    elif cidade == "Belém":
        arquivo_word_origem = os.path.join(diretorio_base, 'Belém.docx')

    if not os.path.exists(arquivo_word_origem):
        msg = QMessageBox()
        msg.setIcon(QMessageBox.Icon.Critical)
        msg.setWindowTitle("Arquivo não encontrado")
        msg.setText(f"Cidade não encontrada.")
        msg.exec()
        return

    if not os.path.exists(arquivo_origem_xlsx):
        if os.path.exists(arquivo_origem_xls):
            converter_xls_para_xlsx(arquivo_origem_xls, arquivo_origem_xlsx)
        else:
            QMessageBox.critical(None, 'Erro', f'O arquivo de origem .xls não foi encontrado: {arquivo_origem_xls}')
            return

    wb_origem = load_workbook(arquivo_origem_xlsx)
    ws_origem = wb_origem.active

    valor_cnpj = str(ws_origem['C14'].value).lower()
    valor_Razao = capitalize_first_letter(str(ws_origem['C7'].value).lower())
    valor_Fantasia = capitalize_first_letter(str(ws_origem['C8'].value).lower())
    valor_Fantasia = f"({valor_Fantasia})"
    valor_Fantasia = valor_Fantasia.replace('#', '')
    valor_Endereco = capitalize_first_letter(str(ws_origem['C9'].value).lower())
    valor_Bairro = capitalize_first_letter(str(ws_origem['C10'].value).lower())
    valor_Cidade = capitalize_first_letter(str(ws_origem['C11'].value).lower())
    valor_Estado = str(ws_origem['N10'].value).upper()
    valor_Cep = str(ws_origem['N12'].value).lower()
    valor_Email = str(ws_origem['C13'].value).lower()
    valor_IE_RG = str(ws_origem['N14'].value).lower()
    valor_Fone = str(ws_origem['C12'].value).lower()
    valor_Num = str(ws_origem['N11'].value).lower()
    valor_Socio = capitalize_first_letter(str(ws_origem['C18'].value).lower())
    valor_Cpf = str(ws_origem['C19'].value).lower()

    match = re.search(r'\((\d+)\)', valor_Razao)
    if match:
        valor_Cod = match.group(1)
        valor_Razao = re.sub(r'\(\d+\)', '', valor_Razao).strip()
    else:
        valor_Cod = ''

    try:
        valor_mensalidade_float = float(valor_mensalidade)
    except ValueError:
        valor_mensalidade_float = 0

    valor_Mensal = formatar_valor_mensalidade(valor_mensalidade)
    valor_Mensal_percentual = "{:.2f}".format(valor_mensalidade_float / 1518 * 100).replace('.', 'X').replace(',', '.').replace('X', ',')

    if valor_adesao == "Isento":
        valor_Adesao = "(ISENTO)"
        valor_Adesao_texto = "ISENTO"
    else:
        valor_Adesao = formatar_valor_adesao(valor_adesao, "de")
        valor_Adesao_texto = num2words(float(valor_adesao), lang='pt_BR', to='currency').upper()
        valor_Adesao = f"{valor_Adesao} ({valor_Adesao_texto})"

    doc = Document(arquivo_word_origem)

    for paragraph in doc.paragraphs:
        replace_placeholder(paragraph, "valor_Razao", valor_Razao, font_size=Pt(9), to_upper=True)
        replace_placeholder(paragraph, "valor_Fantasia", valor_Fantasia, font_size=Pt(9), to_upper=True)
        replace_placeholder(paragraph, "valor_Cidade", valor_Cidade)
        replace_placeholder(paragraph, "valor_Estado", valor_Estado)
        replace_placeholder(paragraph, "valor_Endereco", valor_Endereco)
        replace_placeholder(paragraph, "valor_Num", valor_Num)
        replace_placeholder(paragraph, "valor_Bairro", valor_Bairro)
        replace_placeholder(paragraph, "valor_cnpj", valor_cnpj)
        replace_placeholder(paragraph, "valor_IE_RG", valor_IE_RG)
        replace_placeholder(paragraph, "valor_Socio", valor_Socio)
        replace_placeholder(paragraph, "valor_Cpf", valor_Cpf)
        replace_placeholder(paragraph, "valor_Cod", valor_Cod, font_size=Pt(13))
        replace_placeholder(paragraph, "data_inicio", data_inicio)
        replace_placeholder(paragraph, "data_fim", data_fim)
        replace_placeholder(paragraph, "valor_Mensal", valor_Mensal_percentual)
        replace_placeholder(paragraph, "valor_adesao", valor_Adesao)
        replace_placeholder(paragraph, "data_contrato", data_contrato)
        replace_placeholder(paragraph, "qtd_acesso", qtd_acesso)

    valor_Fantasia = valor_Fantasia.replace("(", "").replace(")", "").strip()
    arquivo_destino_docx = os.path.join(diretorio_base,f'{valor_Cod} - {valor_Fantasia.upper()}.docx')
    doc.save(arquivo_destino_docx)
    QMessageBox.information(None, 'Sucesso', 'Dados inseridos no documento Word com sucesso!')


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("Preenchimento de Documento")
        self.combobox = QComboBox()
        self.combobox.addItems(["Fortaleza", "Belém"])

        self.date_edit_inicio = QDateEdit(calendarPopup=True)
        self.date_edit_inicio.setDate(QDate.currentDate())
        self.date_edit_fim = QDateEdit(calendarPopup=True)
        self.date_edit_fim.setDate(QDate.currentDate())
        self.date_edit_contrato = QDateEdit(calendarPopup=True)
        self.date_edit_contrato.setDate(QDate.currentDate())

        self.line_edit_quantidade = QLineEdit()
        self.line_edit_quantidade.setPlaceholderText('Digite A Quantidade')
        self.line_edit_quantidade.setValidator(QIntValidator(0,9999))


        self.line_edit_mensalidade = QLineEdit()
        self.line_edit_mensalidade.setPlaceholderText("R$ 0,00")
        self.line_edit_mensalidade.setValidator(QDoubleValidator(0.0, 999999.99, 2))
        self.line_edit_mensalidade.editingFinished.connect(self.formatar_mensalidade)

        self.line_edit_adesao = QLineEdit()
        self.line_edit_adesao.setPlaceholderText("R$ 0,00")
        self.line_edit_adesao.setValidator(QDoubleValidator(0.0, 999999.99, 2))
        self.line_edit_adesao.editingFinished.connect(self.formatar_adesao)

        self.checkbox_isento = QCheckBox("Isento")


        self.button_atualizar = QPushButton("Atualizar Documento")

        layout = QVBoxLayout()
        layout.addWidget(QLabel("Cidade:"))
        layout.addWidget(self.combobox)
        layout.addWidget(QLabel("Data de Início:"))
        layout.addWidget(self.date_edit_inicio)
        layout.addWidget(QLabel("Data de Fim:"))
        layout.addWidget(self.date_edit_fim)
        layout.addWidget(QLabel("Data do Contrato:"))
        layout.addWidget(self.date_edit_contrato)
        layout.addWidget(QLabel("Digite A Quantidade De Acessos:"))
        layout.addWidget(self.line_edit_quantidade)
        
        layout.addWidget(QLabel("Valor da Mensalidade:"))
        layout.addWidget(self.line_edit_mensalidade)
        layout.addWidget(QLabel("Valor de Adesão:"))
        layout.addWidget(self.line_edit_adesao)
        layout.addWidget(self.checkbox_isento)
        layout.addWidget(self.button_atualizar)

        container = QWidget()
        container.setLayout(layout)
        self.setCentralWidget(container)

        # Conectar os sinais
        self.button_atualizar.clicked.connect(self.atualizar_documento)
        self.checkbox_isento.toggled.connect(self.alternar_adicao_adesao)

    def formatar_mensalidade(self):
        texto = self.line_edit_mensalidade.text().replace('R$', '').replace('.', '').replace(',', '.')
        texto_formatado = formatar_valor_mensalidade(texto)
        self.line_edit_mensalidade.setText(texto_formatado)

    def formatar_adesao(self):
        texto = self.line_edit_adesao.text().replace('R$', '').replace('.', '').replace(',', '.')
        texto_formatado = formatar_valor_adesao(texto)
        self.line_edit_adesao.setText(texto_formatado)

    def formatar_data_contrato(self, data_qt):

        data = data_qt.toPyDate()

        return data.strftime("%d de %B de %Y")
    

    def atualizar_documento(self):
        data_contrato = self.formatar_data_contrato(self.date_edit_contrato.date())
        data_inicio = self.date_edit_inicio.date().toString('dd/MM/yyyy')
        data_fim = self.date_edit_fim.date().toString('dd/MM/yyyy')
        qtd_acesso = int(self.line_edit_quantidade.text()) if self.line_edit_quantidade.text() else 0
        valor_mensalidade = self.line_edit_mensalidade.text().replace('R$', '').replace('.', '').replace(',', '.')
        valor_adesao = self.line_edit_adesao.text().replace('R$', '').replace('.', '').replace(',', '.')
        if self.checkbox_isento.isChecked():
            valor_adesao = "Isento"
        atualizar_documento(self.combobox.currentText(), data_contrato, data_inicio, data_fim, valor_mensalidade, qtd_acesso, valor_adesao)

    def alternar_adicao_adesao(self):
        if self.checkbox_isento.isChecked():
            self.line_edit_adesao.setText("")
            self.line_edit_adesao.setEnabled(False)
        else:
            self.line_edit_adesao.setText("")
            self.line_edit_adesao.setEnabled(True)

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())